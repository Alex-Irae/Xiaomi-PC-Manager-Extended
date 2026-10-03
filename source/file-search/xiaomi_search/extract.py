"""Read-only extraction and bounded structured chunks. No model or native DLLs."""
import ast
import re
import zipfile
from pathlib import Path

TEXT = {".txt", ".md", ".markdown", ".tex", ".py", ".c", ".cpp", ".h", ".hpp", ".java", ".js", ".jsx", ".ts", ".tsx", ".json", ".yaml", ".yml", ".toml", ".csv", ".log", ".ini", ".rst", ".html", ".css", ".sql", ".ps1", ".sh"}
SUPPORTED = TEXT | {".pdf", ".docx", ".pptx", ".xlsx", ".rtf"}
EXTRACTOR_VERSION = "1"
CHUNK_VERSION = "bounded-structured-1"


def segments(path, maximum):
    """Yield (location, text), preserving pages, slides, headings, and symbols.

    Input is a local Path. Output strings contain plain text, never trusted HTML.
    The cumulative extraction limit prevents very large documents filling memory.
    """
    suffix = path.suffix.lower()
    if suffix in TEXT:
        raw = path.read_bytes()
        if b"\x00" in raw[:4096] and not raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            raise ValueError("Binary content in a text file")
        text = raw.decode("utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig", errors="replace")
        if len(text) > maximum:
            raise ValueError("Extracted text exceeds configured limit")
        if suffix == ".py":
            try:
                tree = ast.parse(text)
                lines = text.splitlines()
                cursor = 0
                for node in tree.body:
                    if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                        start = min([node.lineno] + [d.lineno for d in node.decorator_list]) - 1
                        if start > cursor:
                            yield f"Lines {cursor + 1}-{start}", "\n".join(lines[cursor:start])
                        yield f"{node.name} · line {start + 1}", "\n".join(lines[start:node.end_lineno])
                        cursor = node.end_lineno
                if cursor < len(lines):
                    yield f"Line {cursor + 1}", "\n".join(lines[cursor:])
                return
            except SyntaxError:
                pass
        heading = "Text"
        group = []
        for line in text.splitlines():
            if re.match(r"^#{1,6}\s|^\\(?:sub)*section\*?\{", line) and group:
                yield heading, "\n".join(group)
                group = []
                heading = line.strip()
            group.append(line)
        yield heading, "\n".join(group)
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        total = 0
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            total += len(text)
            if total > maximum:
                raise ValueError("Extracted text exceeds configured limit")
            yield f"Page {number}", text
    elif suffix == ".docx":
        from docx import Document
        heading, group, total = "Document", [], 0
        document = Document(path)
        for paragraph in document.paragraphs:
            total += len(paragraph.text)
            if total > maximum:
                raise ValueError("Extracted text exceeds configured limit")
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                yield heading, "\n".join(group)
                heading, group = paragraph.text, []
            group.append(paragraph.text)
        yield heading, "\n".join(group)
        for number, table in enumerate(document.tables, 1):
            text = "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
            total += len(text)
            if total > maximum:
                raise ValueError("Extracted text exceeds configured limit")
            yield f"Table {number}", text
    elif suffix == ".pptx":
        from pptx import Presentation
        total = 0
        for number, slide in enumerate(Presentation(path).slides, 1):
            texts = [shape.text for shape in slide.shapes if shape.has_text_frame]
            for shape in slide.shapes:
                if shape.has_table:
                    texts += [" | ".join(c.text for c in row.cells) for row in shape.table.rows]
            if slide.has_notes_slide:
                texts.append(slide.notes_slide.notes_text_frame.text)
            text = "\n".join(texts)
            total += len(text)
            if total > maximum:
                raise ValueError("Extracted text exceeds configured limit")
            yield f"Slide {number}", text
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        total = 0
        try:
            for sheet in book:
                header, rows = "", []
                first = 1
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    text = " | ".join(str(value) if value is not None else "" for value in row)
                    total += len(text)
                    if total > maximum:
                        raise ValueError("Extracted text exceeds configured limit")
                    if not header:
                        header = text
                    rows.append(text)
                    if len(rows) >= 20:
                        yield f"{sheet.title} · rows {first}-{number}", header + "\n" + "\n".join(rows)
                        rows, first = [], number + 1
                if rows:
                    yield f"{sheet.title} · row {first}", header + "\n" + "\n".join(rows)
        finally:
            book.close()
    elif suffix == ".rtf":
        from striprtf.striprtf import rtf_to_text
        text = rtf_to_text(path.read_text(encoding="utf-8", errors="replace"))
        if len(text) > maximum:
            raise ValueError("Extracted text exceeds configured limit")
        yield "Document", text


def chunks(path, config):
    """Return bounded (location, text) chunks; overlap stays within each segment."""
    output = []
    if Path(path).suffix.lower() in (".docx", ".pptx", ".xlsx"):
        with zipfile.ZipFile(path) as archive:
            if sum(member.file_size for member in archive.infolist()) > config["max_extracted_characters"] * 16:
                raise ValueError("Office archive exceeds uncompressed safety limit")
    width = config["chunk_characters"]
    overlap = config["overlap_characters"]
    for location, text in segments(Path(path), config["max_extracted_characters"]):
        text = text.strip()
        start = 0
        while start < len(text):
            end = min(start + width, len(text))
            if end < len(text):
                boundary = max(text.rfind("\n", start + width // 2, end), text.rfind(" ", start + width // 2, end))
                if boundary > start:
                    end = boundary
            part = text[start:end].strip()
            if part:
                output.append((location, part))
            if len(output) > config["max_chunks_per_file"]:
                raise ValueError("File exceeds configured chunk limit")
            if end >= len(text):
                break
            start = max(start + 1, end - overlap)
    return output
